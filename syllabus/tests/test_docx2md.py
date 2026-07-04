"""Tests for docx_to_markdown — the .docx -> markdown converter.

Focus: header/footer content reaching both the visible markdown AND the
extract_dates scan (syllabi commonly put the semester/term date in a header,
sometimes laid out as a table rather than a paragraph).

Run from the syllabus/ folder:  .venv/bin/python -m unittest discover tests
"""

import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches

from syllabus.dates import extract_dates
from syllabus.docx2md import docx_to_markdown

TMP = Path("/tmp/syllabus_docx2md_test")
TMP.mkdir(exist_ok=True)


class HeaderFooterParagraphTest(unittest.TestCase):
    def test_header_and_footer_paragraph_dates_reach_output(self):
        doc = Document()
        doc.sections[0].header.paragraphs[0].text = "Header date: March 3, 2026"
        doc.sections[0].footer.paragraphs[0].text = "Footer: see Dec 15"
        doc.add_paragraph("Body text here.")
        path = TMP / "hf_paragraph.docx"
        doc.save(path)

        md = docx_to_markdown(path)
        self.assertIn("March 3, 2026", md)
        self.assertIn("Dec 15", md)

        entries = extract_dates(md, 2026)
        dates = {d for d, _ in entries}
        from datetime import date
        self.assertIn(date(2026, 3, 3), dates)
        self.assertIn(date(2026, 12, 15), dates)


class HeaderTableTest(unittest.TestCase):
    def test_header_table_date_reaches_output(self):
        # School templates commonly lay out header content (course | term |
        # date) as a table rather than paragraphs.
        doc = Document()
        tbl = doc.sections[0].header.add_table(rows=1, cols=3, width=Inches(6))
        tbl.rows[0].cells[0].text = "PSYX362"
        tbl.rows[0].cells[1].text = "Spring 2026"
        tbl.rows[0].cells[2].text = "Updated Jan 15"
        doc.add_paragraph("Body text.")
        path = TMP / "hf_table.docx"
        doc.save(path)

        md = docx_to_markdown(path)
        self.assertIn("Updated Jan 15", md)

        entries = extract_dates(md, 2026)
        from datetime import date
        self.assertIn(date(2026, 1, 15), {d for d, _ in entries})

    def test_no_header_footer_content_is_a_noop(self):
        doc = Document()
        doc.add_paragraph("Plain body, no header/footer content.")
        path = TMP / "no_hf.docx"
        doc.save(path)
        md = docx_to_markdown(path)
        self.assertNotIn("Header/footer", md)




class TrackedChangesTest(unittest.TestCase):
    def test_insertion_text_is_included(self):
        # A run wrapped in w:ins (an insertion) is invisible to plain
        # Paragraph.text; the fix extracts it via _paragraph_text.
        import lxml.etree as etree
        doc = Document()
        p = doc.add_paragraph("Before ")
        ins = etree.SubElement(p._p, qn("w:ins"))
        ins.set(qn("w:id"), "1")
        ins.set(qn("w:author"), "test")
        r = etree.SubElement(ins, qn("w:r"))
        t = etree.SubElement(r, qn("w:t"))
        t.text = "INSERTED-Feb-13"
        p.add_run(" After")
        path = TMP / "tracked_ins.docx"
        doc.save(path)

        md = docx_to_markdown(path)
        self.assertIn("INSERTED-Feb-13", md, "insertion text dropped")

    def test_deletion_text_is_excluded(self):
        # w:delText (deleted content) must NOT appear — it was removed.
        import lxml.etree as etree
        doc = Document()
        p = doc.add_paragraph("Kept ")
        deln = etree.SubElement(p._p, qn("w:del"))
        deln.set(qn("w:id"), "2")
        deln.set(qn("w:author"), "test")
        r = etree.SubElement(deln, qn("w:r"))
        dt = etree.SubElement(r, qn("w:delText"))
        dt.text = "REMOVED-content"
        path = TMP / "tracked_del.docx"
        doc.save(path)

        md = docx_to_markdown(path)
        self.assertIn("Kept", md)
        self.assertNotIn("REMOVED-content", md, "deleted text leaked into output")


class MergedCellTest(unittest.TestCase):
    def test_merged_header_cell_not_duplicated(self):
        doc = Document()
        t = doc.add_table(rows=2, cols=3)
        t.cell(0, 0).merge(t.cell(0, 2))
        t.cell(0, 0).text = "Merged header Feb 13"
        t.cell(1, 0).text = "a"
        t.cell(1, 1).text = "b"
        t.cell(1, 2).text = "c"
        path = TMP / "merged.docx"
        doc.save(path)

        md = docx_to_markdown(path)
        self.assertEqual(md.count("Merged header Feb 13"), 1, "merged cell text duplicated")
        # Row still has 3 columns for table validity.
        header_row = md.splitlines()[0]
        self.assertEqual(header_row.count("|"), 4)


class CorruptInputTest(unittest.TestCase):
    def test_non_zip_docx_raises_value_error(self):
        bad = TMP / "corrupt.docx"
        bad.write_bytes(b"not a docx file at all")
        with self.assertRaises(ValueError):
            docx_to_markdown(bad)

    def test_partial_zip_docx_raises_value_error(self):
        bad = TMP / "partial.docx"
        bad.write_bytes(b"PK\x03\x04" + b"\x00" * 30)
        with self.assertRaises(ValueError):
            docx_to_markdown(bad)
if __name__ == "__main__":
    unittest.main()
